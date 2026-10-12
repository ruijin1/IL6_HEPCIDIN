# -*- coding: utf-8 -*-
r"""Generic Markdown manuscript -> submission-ready docx (Times New Roman 12 pt, double spacing, real Word tables)

Usage: python 20261009_md2docx_generic.py <input.md> <output.docx> [--numbered]

Optional switches:
  --numbered  centre an Arabic page number in the footer and enable **continuous line numbering** (the
              Scientific Reports guidelines recommend that submitted manuscripts carry page/line numbers
              so that reviewers can comment by position).
              Without this switch everything stays as is (no page number, no line numbers).

Generalised from 20261007_md2docx_manuscript.py, with additions:
  - 【待补：…】/【TODO…】 markers are rendered in **bold red** so they are easy to spot in Word
  - the first table row is bold + 9 pt + single spacing (so table notes stay readable)
  - the top "> Nature: ..." notes block is rendered as small single-spaced text and can be deleted wholesale before submission
  - --numbered: footer page number (PAGE field) + continuous line numbers (w:lnNumType)
"""
import sys, os, re
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if len(sys.argv) < 3:
    print("usage: python 20261009_md2docx_generic.py <input.md> <output.docx> [--numbered]"); sys.exit(1)
SRC, DST = sys.argv[1], sys.argv[2]
NUMBERED = "--numbered" in sys.argv[3:]
T = open(SRC, encoding="utf-8").read()
lines = T.split("\n")

def _add_page_number_footer(section):
    """Insert a centred PAGE field (Arabic page number) into the footer"""
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run()
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)

def _enable_line_numbers(section):
    """Enable continuous line numbering (w:lnNumType must sit after pgMar and before cols)"""
    sectPr = section._sectPr
    if sectPr.find(qn("w:lnNumType")) is not None:
        return
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")          # spacing between the line number and the text (twips)
    anchor = None
    for tag in ("w:cols", "w:pgNumType", "w:docGrid"):
        found = sectPr.find(qn(tag))
        if found is not None:
            anchor = found
            break
    if anchor is not None:
        anchor.addprevious(ln)
    else:
        sectPr.append(ln)

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin = s.bottom_margin = Cm(2.5)
    if NUMBERED:
        _add_page_number_footer(s)
        _enable_line_numbers(s)


st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
st.paragraph_format.space_after = Pt(0)

TODO = re.compile(r'(【待补[^】]*】|【TODO[^】]*】)')

def _unesc(s):
    return s.replace("\x01", "*").replace("\x02", "_")

def _style(run):
    run.font.name = "Times New Roman"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def add_runs(par, text, size=None):
    """Handle **bold** / *italic*, and mark 【待补…】 in bold red"""
    text = text.replace("\\*", "\x01").replace("\\_", "\x02")
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|【待补[^】]*】|【TODO[^】]*】)', text):
        if not tok:
            continue
        if TODO.fullmatch(tok):
            r = par.add_run(tok); r.bold = True; r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        elif tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(_unesc(tok[2:-2])); r.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            r = par.add_run(_unesc(tok[1:-1])); r.italic = True
        else:
            r = par.add_run(_unesc(tok))
        _style(r)
        if size:
            r.font.size = Pt(size)

def heading(text, level):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(14 if level == 1 else (11 if level == 2 else 8))
    p.paragraph_format.space_after = Pt(4)
    add_runs(p, text, size=(15 if level == 1 else (13 if level == 2 else 12)))
    for r in p.runs:
        r.bold = True
    return p

i, n_tables, n_todo = 0, 0, 0
while i < len(lines):
    ln = lines[i].rstrip()
    if not ln.strip():
        i += 1; continue
    if re.fullmatch(r'-{3,}', ln.strip()):
        i += 1; continue

    m = re.match(r'^(#{1,4})\s+(.*)$', ln)
    if m:
        heading(m.group(2).strip(), len(m.group(1))); i += 1; continue

    # Table
    if ln.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r'^\s*\|[\s:|-]+\|\s*$', lines[i + 1]):
        header = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows, j = [], i + 2
        while j < len(lines) and lines[j].lstrip().startswith("|"):
            rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")]); j += 1
        tb = doc.add_table(rows=1, cols=len(header))
        tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, h in enumerate(header):
            c = tb.rows[0].cells[k]; c.text = ""
            add_runs(c.paragraphs[0], h, size=9)
            for r in c.paragraphs[0].runs: r.bold = True
            c.paragraphs[0].paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        for row in rows:
            cells = tb.add_row().cells
            for k in range(len(header)):
                v = row[k] if k < len(row) else ""
                cells[k].text = ""
                add_runs(cells[k].paragraphs[0], v, size=9)
                cells[k].paragraphs[0].paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        doc.add_paragraph(); n_tables += 1; i = j; continue

    # Block quote
    if ln.lstrip().startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        add_runs(p, ln.lstrip("> ").strip(), size=10)
        for r in p.runs: r.italic = True
        i += 1; continue

    # Normal paragraph
    buf, j = [ln], i + 1
    while j < len(lines) and lines[j].strip() and not re.match(r'^(#{1,4}\s|\s*\||>|---)', lines[j]):
        buf.append(lines[j].rstrip()); j += 1
    text = re.sub(r'\s{2,}', ' ', " ".join(x.strip() for x in buf))
    p = doc.add_paragraph()
    add_runs(p, text)
    i = j

n_todo = sum(len(TODO.findall(p.text)) for p in doc.paragraphs) + \
         sum(len(TODO.findall(c.text)) for t in doc.tables for row in t.rows for c in row.cells)
doc.save(DST)
d = Document(DST)
print("Generated:", DST)
print("  %d paragraphs | %d tables | %d TODO markers (bold red in Word)" % (len(d.paragraphs), len(d.tables), n_todo))
print("  page number/line numbers: %s" % ("enabled (footer page number + continuous line numbers)" if NUMBERED else "disabled (add --numbered to enable)"))
print("  size %.1f KB" % (os.path.getsize(DST) / 1024))
